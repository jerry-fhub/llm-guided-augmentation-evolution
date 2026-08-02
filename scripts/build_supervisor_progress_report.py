from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_MD = ROOT / "supervisor_progress_report_formal.md"
OUT_DOCX = ROOT / "supervisor_progress_report_formal.docx"

FIG_DIR = ROOT / "results" / "combined" / "completed_experiment_visuals"


META = {
    "Student": "Yushun Feng",
    "Student ID": "psxyf6",
    "Module": "COMP4026 Research Project in Computer Science (Artificial Intelligence)",
    "Supervisor": "Dr. Chao Chen",
    "Date": "25 July 2026",
}


def rel(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9E2EC", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, size=None, bold=None, italic=None, color=None) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def paragraph(doc, text="", style=None, before=0, after=6, align=None, bold=False, italic=False, color=None, size=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0], bold=True, color="1F4E79", size=16 if level == 1 else 13)
    else:
        r = p.add_run(text)
        set_run_font(r, bold=True, color="1F4E79", size=16 if level == 1 else 13)
    return p


def add_callout(doc, title: str, body: str, fill="F2F6FA") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_width(cell, 9200)
    set_cell_margins(cell, top=120, start=180, bottom=120, end=180)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, bold=True, color="1F4E79", size=10.5)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5)
    set_table_borders(table, color="C9D6E2")
    paragraph(doc, "", after=2)


def add_table(doc, headers, rows, widths, font_size=9.2) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_width(hdr[i], widths[i])
        set_cell_margins(hdr[i])
        set_cell_shading(hdr[i], "E8EEF5")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_run_font(r, bold=True, color="1F3A5F", size=font_size)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_width(cells[i], widths[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
            if i >= len(row) - 2 and str(value).replace(".", "").replace("%", "").replace("-", "").isdigit():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc, "", after=4)


def add_bullets(doc, items) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, size=10.5)


def add_figure(doc, path: Path, caption: str, width=6.25) -> None:
    if not path.exists():
        paragraph(doc, f"[Missing figure: {path.name}]", color="9B1C1C")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    cap = paragraph(doc, caption, after=8, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, color="555555", size=9)


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, color in [
        ("Heading 1", 16, "1F4E79"),
        ("Heading 2", 13, "1F4E79"),
        ("Heading 3", 11.5, "1F3A5F"),
    ]:
        st = styles[style_name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)

    header = section.header.paragraphs[0]
    header.text = "MSc Project Progress Report"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if header.runs:
        set_run_font(header.runs[0], color="666666", size=8.5)
    footer = section.footer.paragraphs[0]
    footer.text = "Yushun Feng | COMP4026"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if footer.runs:
        set_run_font(footer.runs[0], color="666666", size=8.5)


def build_markdown() -> str:
    dashboard = rel(FIG_DIR / "supervisor_result_dashboard.png")
    pipeline = rel(FIG_DIR / "supervisor_method_pipeline.png")
    sota = rel(FIG_DIR / "sota_context_comparison.png")
    gantt = rel(FIG_DIR / "updated_project_gantt.png")

    return f"""# MSc Project Progress Report

**Project title:** LLM-Guided Evolutionary Search for Data Augmentation Policy Optimization in Low-Data Image Recognition  
**Student:** {META["Student"]}  
**Student ID:** {META["Student ID"]}  
**Module:** {META["Module"]}  
**Supervisor:** {META["Supervisor"]}  
**Date:** {META["Date"]}

## Executive Summary

This report summarises the current progress of my MSc dissertation project. The project has evolved from an initial idea of asking an LLM to generate image augmentation policies into a more controlled research framework: the LLM is used as a semantic mutation and crossover operator inside an evaluator-controlled evolutionary search loop. Candidate policies are represented as constrained JSON objects, checked by a validator, optionally repaired, evaluated through supervised or FixMatch training, and selected through ranked feedback.

The current evidence supports a focused research claim. Unconstrained or weakly constrained LLM augmentation search is unstable. Baseline-seeded ranked LLM evolution is more reliable: on CIFAR-10 supervised low-data classification, the best OpenAI-evolved policy almost matches the strongest local Mixup baseline in test accuracy and achieves a slightly higher macro-F1. In EuroSAT, the results show that conservative or no augmentation is preferable under the current setting. In FixMatch, semi-supervised learning improves the absolute CIFAR-10 accuracy substantially, and the in-loop LLM evolution prototype successfully generates valid strong-branch policies after repair. The current LLM-evolved FixMatch policies remain below RandAugment and TrivialAugment, so the next phase should focus on multi-seed evaluation, ablation studies, and a stronger FixMatch-in-the-loop search budget.

![Current evidence dashboard]({dashboard})

## Current Research Question

The project now focuses on the following research question:

**Can baseline-seeded ranked LLM evolution generate valid, interpretable, and competitive data augmentation policies for low-data image recognition, especially when the search is placed inside the FixMatch strong-augmentation branch?**

This question keeps the project aligned with the original computer vision topic while giving the method a clearer research contribution. The emphasis is on a reproducible search framework, controlled LLM policy generation, policy validation and repair, and empirical analysis against strong local baselines.

## Methodology Overview

The method is an evaluator-controlled LLM evolution loop. The LLM proposes candidate augmentation policies, and acceptance is determined by the local training evaluator.

![Method pipeline]({pipeline})

The pipeline has five main components. First, augmentation policies are encoded as JSON genotypes containing operations, probabilities, magnitudes, and optional Mixup/CutMix parameters. Second, all policies pass through a validator that checks operation names, parameter ranges, sub-policy length, and dataset-specific warnings. Third, the initial population is seeded with strong conventional baselines such as standard augmentation, Mixup, CutMix, RandAugment, TrivialAugment, and no augmentation. Fourth, the LLM receives ranked population feedback and generates mutation or crossover children. Fifth, rough-to-full evaluation screens candidates cheaply before selected policies receive fuller evaluation.

The latest method extends this loop to FixMatch. In this setting, weak augmentation produces pseudo-labels and strong augmentation is used for consistency training. This makes the strong augmentation branch a natural target for LLM-guided policy evolution.

## Experiment Results

### Supervised Experiments

| Dataset | Best local baseline | Best LLM/evolution result | Interpretation |
|---|---:|---:|---|
| CIFAR-10 early supervised run | RandAugment 43.97% | One-shot LLM 45.25%; LLM evolution 40.48% | Early LLM search was runnable but unstable. |
| Flowers102 early supervised run | TrivialAugment 47.47% | LLM evolution 31.86% | Conventional augmentation remained stronger. |
| EuroSAT early supervised run | Standard 75.75% | LLM evolution 68.68% | Strong augmentation can damage domain-specific signals. |
| CIFAR-10 strong OpenAI run | Mixup 48.50% | OpenAI-evolved `gen02_mut_005` 48.45% | Baseline-seeded ranked LLM evolution became competitive. |
| EuroSAT strong OpenAI run | No augmentation 73.95% | OpenAI-evolved `gen01_mut_002` 72.45% | Conservative policies are preferred in this setting. |

The strongest supervised result is on CIFAR-10. The best OpenAI-evolved policy keeps RandomCrop, HorizontalFlip, light ColorJitter, and Mixup with `mixup_alpha=0.2`. This is a meaningful result because the LLM preserves the strongest local baseline prior and makes a small, interpretable modification while avoiding an overly complex policy.

### FixMatch Experiments

| Method | Test accuracy | Test macro-F1 | Validation accuracy | Interpretation |
|---|---:|---:|---:|---|
| Supervised Mixup baseline | 48.50% | 46.30% | 47.40% | Best supervised local accuracy baseline. |
| Supervised LLM best | 48.45% | 47.22% | 50.20% | Best supervised LLM-evolved policy. |
| FixMatch standard | 56.75% | 55.12% | 56.60% | Unlabelled data improves performance. |
| FixMatch reused LLM policy | 59.00% | 58.38% | 61.40% | Supervised LLM policy transfers usefully to FixMatch. |
| FixMatch repaired LLM child `mut_001` | 59.05% | 58.05% | 62.70% | In-loop LLM evolution is functional, with a very small gain over reuse. |
| FixMatch TrivialAugment | 61.50% | 61.08% | 63.00% | Strong canonical augmentation baseline. |
| FixMatch RandAugment | 62.05% | 61.05% | 63.90% | Current best local result. |

The FixMatch experiments confirm that semi-supervised learning is a better setting for studying strong augmentation policies. The best local accuracy rises from 48.50% in supervised training to 62.05% with FixMatch RandAugment. The best in-loop LLM child reaches 59.05%, which is useful and valid but still behind canonical strong augmentation.

## SOTA and Strong-Baseline Context

![SOTA context comparison]({sota})

External papers provide an important reference point. UDA reports 94.57% CIFAR-10 accuracy in a 250-label semi-supervised protocol, and FixMatch reports 94.93% under a similar protocol. Full-data automated augmentation methods such as AutoAugment and Population Based Augmentation report approximately 98.5% CIFAR-10 accuracy under much larger training and search budgets. These numbers are not directly comparable to my local 15-epoch, ResNet18-CIFAR, student-scale experiments. They show the performance gap and help define future improvement targets.

The dissertation should therefore make local strong-baseline claims and use SOTA papers as external context. The current contribution is the implementation and analysis of a controlled LLM-guided augmentation policy evolution framework, including baseline seeding, ranked LLM feedback, validation, repair, rough-to-full evaluation, and FixMatch-aware strong-branch policy search.

## Completed Work and Remaining Work

| Work item | Status | Evidence | Remaining action |
|---|---|---|---|
| Literature review and research narrowing | Completed, needs final update | Introduction/background and reference library updated | Add final experimental interpretation to Discussion. |
| Dataset preparation and low-data splits | Completed | CIFAR-10, Flowers102, EuroSAT loaders and deterministic splits | Keep CIFAR-10 and EuroSAT as core final datasets. |
| Supervised baseline pipeline | Completed | Baseline result tables and combined summaries | Add selected multi-seed reruns. |
| JSON policy schema, transform builder, validator | Completed | Policy, builder, and validator modules | Present as core engineering contribution. |
| Baseline-seeded ranked OpenAI evolution | Completed | CIFAR-10 and EuroSAT strong OpenAI results | Analyse policy lineage and operation frequency. |
| FixMatch training and policy reuse | Completed | FixMatch standard, RandAugment, TrivialAugment, and LLM policy results | Add per-class diagnostics. |
| FixMatch-in-the-loop LLM evolution | Prototype completed | Repaired LLM children evaluated under FixMatch | Extend to 2-3 generations and larger offspring pool. |
| Policy repair layer | Prototype completed | Invalid LLM children repaired and evaluated | Add repair-rate and repair-impact analysis. |
| SOTA/strong baseline comparison | Completed for progress report | SOTA context figure and verified references | Convert into thesis discussion subsection. |
| Ablation studies | Not completed | No complete ablation table yet | Run baseline seeding, ranked feedback, repair, and rough/full ablations. |
| Statistical reliability | Not completed | Most strong runs are single seed | Run 3-seed comparisons for final policies and baselines. |
| Thesis Methodology | Draft completed | Methodology LaTeX chapter created | Revise after supervisor feedback. |
| Results, Discussion, Conclusion chapters | Not completed | Existing reports and figures available | Write final chapters after additional experiments. |

## Updated Project Plan

![Updated Gantt chart]({gantt})

The next phase should prioritise reliability and dissertation readiness. The immediate plan is to run multi-seed final comparisons, add ablations, extend the FixMatch-in-loop search, and then convert the current experiment reports into the Results and Discussion chapters.

## Risks and Mitigation

The main risk is statistical reliability. Several key experiments are single-seed, and the difference between the best supervised LLM policy and Mixup is extremely small. I will mitigate this by running at least three seeds for the final CIFAR-10 FixMatch baselines and the best LLM-evolved policies.

The second risk is rough-to-full ranking noise. Rough evaluation helps filter weak policies, but it does not always predict final performance reliably, especially in FixMatch. I will analyse rough/full correlation and consider repeated rough seeds for shortlisted candidates.

The third risk is search-space mismatch. The current JSON policy space approximates RandAugment-like behaviour, while torchvision RandAugment and TrivialAugment are stronger canonical implementations. I will consider adding RandAugment or TrivialAugment as macro-operations in the search space.

## Questions for Supervisor Discussion

1. Is the current contribution framing appropriate: evaluator-controlled LLM augmentation policy evolution with direct SOTA performance treated as future work?
2. Should the final thesis prioritise CIFAR-10 + EuroSAT, with Flowers102 kept as an early boundary result?
3. Which additional experiment is most valuable for the dissertation: multi-seed reliability, ablation studies, or extended FixMatch-in-the-loop evolution?
4. Is the proposed Methodology chapter scope sufficient, or should the FixMatch component be treated as the central method?

## Key References

- FixMatch: https://arxiv.org/abs/2001.07685
- UDA: https://arxiv.org/abs/1904.12848
- ReMixMatch: https://arxiv.org/abs/1911.09785
- AutoAugment: https://arxiv.org/abs/1805.09501
- Population Based Augmentation: https://proceedings.mlr.press/v97/ho19b.html
- RandAugment: https://papers.nips.cc/paper/2020/hash/d85b63ef0ccb114d0a3bb7b7d808028f-Abstract.html
- TrivialAugment: https://openaccess.thecvf.com/content/ICCV2021/html/Muller_TrivialAugment_Tuning-Free_Yet_State-of-the-Art_Data_Augmentation_ICCV_2021_paper.html
- FunSearch: https://www.nature.com/articles/s41586-023-06924-6
- LLM feedback augmentation policy optimisation: https://arxiv.org/abs/2410.13453

## Appendix: Key Artefacts

| Artefact | Path |
|---|---|
| Main experiment entry point | `experiments/run_experiment.py` |
| Supervised evolution loop | `src/image_aug_evolution/search/evolution.py` |
| FixMatch evolution loop | `src/image_aug_evolution/search/fixmatch_evolution.py` |
| LLM policy generator | `src/image_aug_evolution/llm/openai_generator.py` |
| Policy validator | `src/image_aug_evolution/augmentation/validator.py` |
| FixMatch trainer | `src/image_aug_evolution/models/fixmatch.py` |
| Combined result summary | `results/combined/all_method_summaries.csv` |
| Best supervised CIFAR-10 LLM policy | `results/cifar10_resnet18cifar_strong_openai/policies/gen02_mut_005.json` |
| Best repaired FixMatch LLM child | `results/cifar10_fixmatch_repaired_llm_children/policies/fm_gen01_mut_001_repaired.json` |
"""


def build_docx() -> None:
    doc = Document()
    configure_doc(doc)

    paragraph(doc, "MSc Project Progress Report", size=24, bold=True, color="1F4E79", after=4)
    paragraph(doc, "LLM-Guided Evolutionary Search for Data Augmentation Policy Optimization in Low-Data Image Recognition", size=13.5, color="444444", after=14)
    for key, value in META.items():
        p = paragraph(doc, "", after=2)
        r1 = p.add_run(f"{key}: ")
        set_run_font(r1, bold=True, size=10.5)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.5)

    add_callout(
        doc,
        "Current position",
        "The project has a working evaluator-controlled LLM evolution framework. The strongest local result remains FixMatch with RandAugment, while the LLM-evolved FixMatch child is valid and competitive but still below canonical strong augmentation. The thesis should emphasise framework design, mechanism analysis, controlled comparison, and next-step reliability experiments.",
        fill="EAF3F8",
    )

    add_figure(doc, FIG_DIR / "supervisor_result_dashboard.png", "Figure 1. Current experiment evidence dashboard.", width=6.05)

    heading(doc, "1. Executive Summary")
    paragraph(doc, "This report summarises the current progress of my MSc dissertation project. The project has evolved from an initial idea of asking an LLM to generate image augmentation policies into a controlled research framework: the LLM is used as a semantic mutation and crossover operator inside an evaluator-controlled evolutionary search loop. Candidate policies are represented as constrained JSON objects, checked by a validator, optionally repaired, evaluated through supervised or FixMatch training, and selected through ranked feedback.")
    paragraph(doc, "The current evidence supports a focused research claim. Unconstrained or weakly constrained LLM augmentation search is unstable. Baseline-seeded ranked LLM evolution is more reliable: on CIFAR-10 supervised low-data classification, the best OpenAI-evolved policy almost matches the strongest local Mixup baseline in test accuracy and achieves a slightly higher macro-F1. In EuroSAT, the results show that conservative or no augmentation is preferable under the current setting. In FixMatch, semi-supervised learning improves CIFAR-10 accuracy substantially, and the in-loop LLM evolution prototype successfully generates valid strong-branch policies after repair.")

    heading(doc, "2. Current Research Question")
    add_callout(
        doc,
        "Research question",
        "Can baseline-seeded ranked LLM evolution generate valid, interpretable, and competitive data augmentation policies for low-data image recognition, especially when the search is placed inside the FixMatch strong-augmentation branch?",
        fill="F4F6F9",
    )
    paragraph(doc, "This question keeps the project aligned with computer vision while clarifying the research contribution. The emphasis is on reproducible search, controlled LLM policy generation, policy validation and repair, and empirical analysis against strong local baselines.")

    heading(doc, "3. Methodology Overview")
    paragraph(doc, "The method is an evaluator-controlled LLM evolution loop. The LLM proposes candidate augmentation policies, and acceptance depends on local validation and test performance. The latest implementation extends the loop to FixMatch, where weak augmentation produces pseudo-labels and strong augmentation is used for consistency training.")
    add_figure(doc, FIG_DIR / "supervisor_method_pipeline.png", "Figure 2. Evaluator-controlled LLM evolution pipeline.", width=6.05)
    add_bullets(doc, [
        "Policy representation: augmentation policies are JSON genotypes with operations, probabilities, magnitudes, and optional Mixup/CutMix parameters.",
        "Constraint handling: every policy passes through a validator, and FixMatch children can be repaired when they exceed local search-space constraints.",
        "Baseline seeding: initial populations include standard augmentation, Mixup, CutMix, RandAugment, TrivialAugment, and no augmentation.",
        "Ranked LLM feedback: the LLM receives policy rankings, validation metrics, objective scores, complexity notes, and baseline references.",
        "Rough-to-full evaluation: shorter training filters candidates before selected policies receive fuller evaluation.",
    ])

    heading(doc, "4. Main Experimental Results")
    paragraph(doc, "The table below summarises the main supervised results. Early experiments established that weak LLM search was runnable but unstable. The stronger OpenAI-ranked implementation is the main supervised evidence.")
    add_table(
        doc,
        ["Dataset / setting", "Best local baseline", "Best LLM/evolution result", "Interpretation"],
        [
            ["CIFAR-10 early supervised", "RandAugment 43.97%", "One-shot LLM 45.25%; LLM evolution 40.48%", "Early LLM search was unstable."],
            ["Flowers102 early supervised", "TrivialAugment 47.47%", "LLM evolution 31.86%", "Conventional augmentation remained stronger."],
            ["EuroSAT early supervised", "Standard 75.75%", "LLM evolution 68.68%", "Strong augmentation can damage domain-specific signals."],
            ["CIFAR-10 strong OpenAI", "Mixup 48.50%", "OpenAI-evolved gen02_mut_005 48.45%", "Baseline-seeded ranked LLM evolution became competitive."],
            ["EuroSAT strong OpenAI", "No augmentation 73.95%", "OpenAI-evolved gen01_mut_002 72.45%", "Conservative policies are preferred."],
        ],
        [1900, 1850, 2450, 3040],
        font_size=8.8,
    )

    paragraph(doc, "The FixMatch experiments are now the most relevant evidence because strong augmentation is central to FixMatch's pseudo-label consistency objective.")
    add_table(
        doc,
        ["Method", "Test accuracy", "Macro-F1", "Validation accuracy", "Interpretation"],
        [
            ["Supervised Mixup baseline", "48.50%", "46.30%", "47.40%", "Best supervised local accuracy baseline."],
            ["Supervised LLM best", "48.45%", "47.22%", "50.20%", "Best supervised LLM-evolved policy."],
            ["FixMatch standard", "56.75%", "55.12%", "56.60%", "Unlabelled data improves performance."],
            ["FixMatch reused LLM policy", "59.00%", "58.38%", "61.40%", "Supervised LLM policy transfers usefully to FixMatch."],
            ["FixMatch repaired LLM child mut_001", "59.05%", "58.05%", "62.70%", "In-loop LLM evolution is functional, with a very small gain over reuse."],
            ["FixMatch TrivialAugment", "61.50%", "61.08%", "63.00%", "Strong canonical augmentation baseline."],
            ["FixMatch RandAugment", "62.05%", "61.05%", "63.90%", "Current best local result."],
        ],
        [2500, 1150, 1050, 1350, 3310],
        font_size=8.4,
    )

    heading(doc, "5. SOTA and Strong-Baseline Context")
    paragraph(doc, "External papers provide an important reference point. UDA reports 94.57% CIFAR-10 accuracy in a 250-label semi-supervised protocol, and FixMatch reports 94.93% under a similar protocol. Full-data automated augmentation methods such as AutoAugment and Population Based Augmentation report approximately 98.5% CIFAR-10 accuracy under much larger training and search budgets. These numbers frame the gap between local student-scale experiments and paper-level SOTA.")
    add_figure(doc, FIG_DIR / "sota_context_comparison.png", "Figure 3. Scope-aware comparison between local results and external SOTA context.", width=6.05)
    paragraph(doc, "The dissertation should make local strong-baseline claims and use SOTA papers as external context. The current contribution is the implementation and analysis of a controlled LLM-guided augmentation policy evolution framework, including baseline seeding, ranked LLM feedback, validation, repair, rough-to-full evaluation, and FixMatch-aware strong-branch policy search.")

    heading(doc, "6. Completed Work and Remaining Work")
    add_table(
        doc,
        ["Work item", "Status", "Remaining action"],
        [
            ["Literature review and research narrowing", "Completed, final update needed", "Add final experimental interpretation to Discussion."],
            ["Dataset preparation and low-data splits", "Completed", "Keep CIFAR-10 and EuroSAT as core final datasets."],
            ["Supervised baseline and policy pipeline", "Completed", "Add selected multi-seed reruns."],
            ["Baseline-seeded ranked OpenAI evolution", "Completed", "Analyse policy lineage and operation frequency."],
            ["FixMatch training and policy reuse", "Completed", "Add per-class diagnostics."],
            ["FixMatch-in-the-loop LLM evolution", "Prototype completed", "Extend to 2-3 generations and larger offspring pool."],
            ["Policy repair layer", "Prototype completed", "Add repair-rate and repair-impact analysis."],
            ["Ablation studies", "Not completed", "Run baseline seeding, ranked feedback, repair, and rough/full ablations."],
            ["Statistical reliability", "Not completed", "Run 3-seed comparisons for final policies and baselines."],
            ["Thesis Results, Discussion, Conclusion", "Not completed", "Write final chapters after additional experiments."],
        ],
        [3350, 2300, 3710],
        font_size=8.7,
    )

    heading(doc, "7. Updated Project Plan")
    add_figure(doc, FIG_DIR / "updated_project_gantt.png", "Figure 4. Updated Gantt chart for experiments and dissertation writing.", width=6.05)
    paragraph(doc, "The next phase prioritises reliability and dissertation readiness: multi-seed comparisons, ablation studies, extended FixMatch-in-loop search, policy diagnostics, and conversion of experiment reports into thesis chapters.")

    heading(doc, "8. Risks and Mitigation")
    add_table(
        doc,
        ["Risk", "Impact", "Mitigation"],
        [
            ["Single-seed evidence", "Small differences cannot support strong statistical claims.", "Run 3-seed comparisons for final baselines and LLM policies."],
            ["Rough-to-full ranking noise", "Rough evaluation may select policies that do not remain strong under full training.", "Analyse rough/full correlation and use repeated rough seeds for shortlisted policies."],
            ["Search-space mismatch", "JSON policies may approximate RandAugment less strongly than torchvision implementations.", "Add RandAugment or TrivialAugment as macro-operations."],
            ["Invalid LLM outputs", "Generated policies may exceed local constraints.", "Use validator and repair layer; report invalid-to-valid repair rate."],
        ],
        [2400, 3350, 3610],
        font_size=8.7,
    )

    heading(doc, "9. Questions for Supervisor Discussion")
    add_bullets(doc, [
        "Is the contribution framing appropriate: evaluator-controlled LLM augmentation policy evolution with strong local baselines?",
        "Should the final thesis prioritise CIFAR-10 and EuroSAT, with Flowers102 kept as an early boundary result?",
        "Which additional experiment is most valuable: multi-seed reliability, ablation studies, or extended FixMatch-in-the-loop evolution?",
        "Should FixMatch be treated as the central method in the final dissertation structure?",
    ])

    heading(doc, "Appendix A. Key References")
    add_bullets(doc, [
        "FixMatch: https://arxiv.org/abs/2001.07685",
        "UDA: https://arxiv.org/abs/1904.12848",
        "ReMixMatch: https://arxiv.org/abs/1911.09785",
        "AutoAugment: https://arxiv.org/abs/1805.09501",
        "Population Based Augmentation: https://proceedings.mlr.press/v97/ho19b.html",
        "RandAugment: https://papers.nips.cc/paper/2020/hash/d85b63ef0ccb114d0a3bb7b7d808028f-Abstract.html",
        "TrivialAugment: https://openaccess.thecvf.com/content/ICCV2021/html/Muller_TrivialAugment_Tuning-Free_Yet_State-of-the-Art_Data_Augmentation_ICCV_2021_paper.html",
        "FunSearch: https://www.nature.com/articles/s41586-023-06924-6",
        "LLM feedback augmentation policy optimisation: https://arxiv.org/abs/2410.13453",
    ])

    heading(doc, "Appendix B. Key Artefacts")
    add_table(
        doc,
        ["Artefact", "Path"],
        [
            ["Main experiment entry point", "experiments/run_experiment.py"],
            ["Supervised evolution loop", "src/image_aug_evolution/search/evolution.py"],
            ["FixMatch evolution loop", "src/image_aug_evolution/search/fixmatch_evolution.py"],
            ["LLM policy generator", "src/image_aug_evolution/llm/openai_generator.py"],
            ["Policy validator", "src/image_aug_evolution/augmentation/validator.py"],
            ["FixMatch trainer", "src/image_aug_evolution/models/fixmatch.py"],
            ["Combined result summary", "results/combined/all_method_summaries.csv"],
            ["Best supervised CIFAR-10 LLM policy", "results/cifar10_resnet18cifar_strong_openai/policies/gen02_mut_005.json"],
            ["Best repaired FixMatch LLM child", "results/cifar10_fixmatch_repaired_llm_children/policies/fm_gen01_mut_001_repaired.json"],
        ],
        [3000, 6360],
        font_size=8.5,
    )

    doc.save(OUT_DOCX)


def main() -> None:
    OUT_MD.write_text(build_markdown(), encoding="utf-8")
    build_docx()
    print(f"Wrote {OUT_MD}")
    print(f"Wrote {OUT_DOCX}")


if __name__ == "__main__":
    main()
